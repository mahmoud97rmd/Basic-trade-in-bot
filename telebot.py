"""
Gold Scalper Bot — v4.6 (Deriv WS + MA Gap Filter Edition)
Strategies : STOCH-NEW  |  STOCH-OLD
"""

import asyncio
import aiohttp
import websockets
import json
import os
import pandas as pd
import numpy as np
from datetime import datetime, timedelta, timezone
from metaapi_cloud_sdk import MetaApi
from aiohttp import web

# ─────────────────────────────────────────────────────────────
# CONFIGURATION
# ─────────────────────────────────────────────────────────────
METAAPI_TOKEN = 'eyJhbGciOiJSUzUxMiIsInR5cCI6IkpXVCJ9..NRMo-BO9ezZBEb4XmCQzkMsRN1iAz1rVSk7XWFP-ZGS_AZEyxSfIjnJ5w-r4egazV7tnxNLjjMuAdUb25T3ur3XWKCL4Jo9LFPy9tZzhIMRtlhq8d6YAHK9uxJclqJv5BZQFDeMeiFtyalLNjaE100Lp2zEnGWwlloxF-dpCw5DXvVKeGfMyVx4L2kisshcysDo7OeMkDBU1UB7leHi2eviEl7XQCpmhxdzT4BwMkf8YERx2jouKVu8-koVy00aon0drktGBSlQDOFw2WV0hg-VUfeCBR_Hgw2czqKVJ_lj_ZN3EsjWirirpiuXWbtwdD-VPokjKtX1z3ugcSTS1nd2iFIzauUHdOfb7Jl0R6cm8FosVS-4Iu046DiMsrxiAJ4PBywOXQhsFzZiePqmil1w5HHCxrw_78HNR9XcjBETMpHx9W48llIeUOkBVbsKfBP5iYtGSjS52i0QgpvHkfKrtXfbkMT0_9yJFG2kfZJHwJ5BJzWT4aKXto3l6iGe45xe4ZJhYhZX_RkC6dxR2w84M-uY-wlqiv_sxjHNOguSyOx4lfaeoq5H-LuJiWpHAYxEJUQWoQAQ7PObZOXCDWLRc_vP2gcbv1qYxTjD54FHnqhyf-oTGzAkWG5CVQFKpp9jTHQ3pXEYTSgIUTfHDbtoesAY1HG3nHcHbwujnqo0'
ACCOUNT_ID    = '7d54fa6f-eaf7-4637-92a1-e0356ee729f8'
TG_TOKEN      = '8779425898:AAFQgqay6IO89I2Sf98PigL28v9AHCcZPMw'

# إعدادات Deriv WebSocket
DERIV_APP_ID  = '1089'
DERIV_WS_URL  = f'wss://ws.binaryws.com/websockets/v3?app_id={DERIV_APP_ID}'
DERIV_SYMBOL  = 'frxXAUUSD'

# جميع الفريمات المتاحة
_TFS = ['1m', '2m', '3m', '4m', '5m', '6m', '10m', '12m', '15m', '20m', '30m', '1h', '2h']

# Daily drawdown limit — 3% of start-of-day balance
DD_LIMIT_PCT = 0.03

# ─────────────────────────────────────────────────────────────
# GLOBAL SHARED HTTP SESSION (For Telegram)
# ─────────────────────────────────────────────────────────────
_http: aiohttp.ClientSession | None = None

def get_http() -> aiohttp.ClientSession:
    global _http
    if _http is None or _http.closed:
        connector = aiohttp.TCPConnector(limit=20, ttl_dns_cache=300)
        timeout   = aiohttp.ClientTimeout(total=30, connect=10)
        _http     = aiohttp.ClientSession(connector=connector, timeout=timeout)
    return _http

def c_log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)

# ─────────────────────────────────────────────────────────────
# GLOBAL STATE & CACHE
# ─────────────────────────────────────────────────────────────
bot_state: dict = {
    'status':            'RUNNING',
    'symbol':            'XAUUSDm',
    'live_connected':    False,
    'timeframes':        _TFS,
    'active_tfs':        {
        '1m': False, '2m': True, '3m': True, '4m': False, '5m': False,
        '6m': False, '10m': False, '12m': False, '15m': False,
        '20m': False, '30m': False, '1h': False, '2h': False
    },
    'cache':             {tf: pd.DataFrame() for tf in _TFS},
    'lot_size':          0.05,
    'pip_value':         0.1,
    'spread_pips':       2.2,
    'chat_id':           None,
    'last_update_id':    0,
    'tp_pips':           {
        '1m': 25, '2m': 30, '3m': 40, '4m': 50, '5m': 70, 
        '6m': 75, '10m': 80, '12m': 85, '15m': 90, 
        '20m': 100, '30m': 120, '1h': 150, '2h': 200
    },
    'sl_pips':           {
        '1m': 100, '2m': 100, '3m': 100, '4m': 100, '5m': 100, 
        '6m': 100, '10m': 120, '12m': 120, '15m': 150, 
        '20m': 150, '30m': 200, '1h': 250, '2h': 300
    },
    'strategy_mode':     'STOCH_OLD',
    'filter_mode':       'NO_MA',
    'stoch_k':           5,
    'stoch_smooth':      5,
    'stoch_d':           5,
    'use_stoch_deep':    True,
    'use_stoch_mid':     True,
    'use_stoch_shal':    False,
    
    # --- MA Gap Filter State (بدلاً من فلتر الثبات) ---
    'use_ma_gap_fixed':  False,
    'ma_gap_pips':       10,
    'use_ma_gap_atr':    False,
    'ma_gap_atr_mult':   0.5,
    
    'use_be':            False,
    'use_atr':           False,
    'use_max_spread':    True,
    'max_spread_pips':   3.0,
    'atr_mult_tp':       1.5,
    'atr_mult_sl':       3.0,
    'tp_tolerance_pips': 5.0,
    'sod_balance':       None,
    'sod_date':          None,
    'dd_triggered':      False,
    'use_danger_filter': True, 
    'market_data':       {tf: 'بانتظار الاتصال (Offline)' for tf in _TFS},
    'last_signal_time':  {tf: None for tf in _TFS},
    'connection_obj':    None,
    'account_obj':       None,
    'is_backtesting':    False,
}

# ─────────────────────────────────────────────────────────────
# TIME FILTER
# ─────────────────────────────────────────────────────────────
_BLOCKED_DAMASCUS_HOURS = {13, 18, 21, 22}

def is_blocked_time(dt_utc: datetime) -> bool:
    damascus_hour = (dt_utc.hour + 3) % 24
    return damascus_hour in _BLOCKED_DAMASCUS_HOURS

def blocked_time_label(dt_utc: datetime) -> str:
    damascus_hour = (dt_utc.hour + 3) % 24
    if damascus_hour in _BLOCKED_DAMASCUS_HOURS:
        return f'منطقة خطر ({damascus_hour}:xx دمشق)'
    return ''

# ─────────────────────────────────────────────────────────────
# DAILY DRAWDOWN PROTECTOR
# ─────────────────────────────────────────────────────────────
async def _capture_sod_balance() -> None:
    if not (bot_state['live_connected'] and bot_state['connection_obj']):
        return
    try:
        info  = await bot_state['connection_obj'].get_account_information()
        bal   = float(info.get('balance', 0))
        today = datetime.now(timezone.utc).strftime('%Y-%m-%d')
        bot_state['sod_balance']  = bal
        bot_state['sod_date']     = today
        bot_state['dd_triggered'] = False
        await send_tg_msg(
            f'📅 <b>بداية يوم جديد</b>\n'
            f'رصيد الافتتاح: <b>${bal:.2f}</b>\n'
            f'حد الخسارة اليومي (3%): <b>${bal * DD_LIMIT_PCT:.2f}</b>\n'
            f'سيتوقف البوت عند الإكويتي: <b>${bal * (1 - DD_LIMIT_PCT):.2f}</b>'
        )
    except Exception as e:
        c_log(f'DD: capture SOD error: {e}')

async def daily_drawdown_monitor() -> None:
    while True:
        await asyncio.sleep(30)
        try:
            if not (bot_state['live_connected'] and bot_state['connection_obj']):
                continue
            today = datetime.now(timezone.utc).strftime('%Y-%m-%d')
            if bot_state['sod_date'] != today or bot_state['sod_balance'] is None:
                await _capture_sod_balance()
                continue
            if bot_state['dd_triggered']: continue

            info     = await bot_state['connection_obj'].get_account_information()
            equity   = float(info.get('equity', 0))
            sod      = bot_state['sod_balance']
            limit    = sod * (1 - DD_LIMIT_PCT)
            used_pct = round((sod - equity) / sod * 100, 2) if sod else 0

            if equity <= limit:
                bot_state['status']       = 'PAUSED'
                bot_state['dd_triggered'] = True
                closed = 0
                try:
                    positions = await bot_state['connection_obj'].get_positions()
                    for p in positions:
                        await bot_state['connection_obj'].close_position(p['id'])
                        closed += 1
                except Exception as ce:
                    c_log(f'DD close positions error: {ce}')

                await send_tg_msg(
                    f'🚨🚨🚨 <b>تم إيقاف البوت تلقائياً</b> 🚨🚨🚨\n\n'
                    f'تم الوصول إلى حد الخسارة اليومية (3%)\n\n'
                    f'💰 رصيد الافتتاح:  <b>${sod:.2f}</b>\n'
                    f'📉 الإكويتي الحالي: <b>${equity:.2f}</b>\n'
                    f'📊 الخسارة:         <b>${sod - equity:.2f}  ({used_pct}%)</b>\n'
                    f'🔒 الحد المسموح:     <b>${sod * DD_LIMIT_PCT:.2f}</b>\n\n'
                    f'{"تم إغلاق " + str(closed) + " صفقة مفتوحة." if closed else "لا توجد صفقات مفتوحة."}\n\n'
                    f'البوت في وضع <b>PAUSED</b>.\n'
                    f'لإعادة التشغيل اضغط RUN من القائمة الرئيسية.\n'
                    f'سيُعاد احتساب الرصيد غداً تلقائياً.',
                    get_main_keyboard(),
                )
        except Exception as e:
            c_log(f'DD monitor error: {e}')

# ─────────────────────────────────────────────────────────────
# BACKTEST PROGRESS TRACKER
# ─────────────────────────────────────────────────────────────
class BtProgress:
    BAR_LEN   = 14
    HEARTBEAT = 15

    def __init__(self, label: str, active_tfs: list, is_advanced: bool = False):
        self.label        = label
        self.active_tfs   = active_tfs
        self.is_advanced  = is_advanced
        self.cancelled    = False
        self.phase        = 'تهيئة...'
        self.tf_done      = 0
        self.tf_total     = len(active_tfs)
        self.current_tf   = ''
        self.bars_done    = 0
        self.bars_total   = 0
        self.win          = 0
        self.loss         = 0
        self.be           = 0
        self.profit       = 0.0
        self.chat_id      = None
        self.msg_id       = None
        self._last_edit   = 0.0
        self._lock        = asyncio.Lock()
        self._hb_task     = None

    def _bar(self, done: int, total: int) -> str:
        if total == 0: return chr(9617) * self.BAR_LEN
        filled = round(done / total * self.BAR_LEN)
        return chr(9608) * filled + chr(9617) * (self.BAR_LEN - filled)

    def _pct(self, done: int, total: int) -> str:
        return f'{round(done / total * 100)}%' if total else '0%'

    def _elapsed(self) -> str:
        secs = int(datetime.now(timezone.utc).timestamp() - self._start_ts)
        m, s = divmod(secs, 60)
        return f'{m}m {s:02d}s'

    def _cancel_kbd(self) -> dict:
        return {'inline_keyboard': [[{'text': 'إيقاف الباك تيست', 'callback_data': 'cancel_bt'}]]}

    def _build_text(self) -> str:
        kind  = 'Advanced' if self.is_advanced else 'Backtest'
        total = self.win + self.loss + self.be
        wr    = f'{round(self.win / total * 100)}%' if total else '-'
        pnl   = f'+${round(self.profit, 2)}' if self.profit >= 0 else f'-${abs(round(self.profit, 2))}'
        icon  = 'UP' if self.profit >= 0 else 'DN'
        if self.bars_total > 0:
            overall = (self.tf_done + self.bars_done / self.bars_total) / max(self.tf_total, 1)
        else:
            overall = self.tf_done / max(self.tf_total, 1)
        ov_bar = self._bar(round(overall * 100), 100)
        ov_pct = f'{round(overall * 100)}%'
        tf_bar = self._bar(self.bars_done, self.bars_total) if self.bars_total else chr(9617) * self.BAR_LEN
        tf_pct = self._pct(self.bars_done, self.bars_total) if self.bars_total else '-'
        lines = [
            f'{kind} - <b>{self.label}</b>',
            f'<b>المرحلة:</b> {self.phase}',
            f'',
            f'<b>التقدم الكلي</b>  {ov_pct}',
            f'<code>[{ov_bar}]</code>',
        ]
        if self.current_tf:
            lines += [
                f'',
                f'<b>الفريم:</b> {self.current_tf}  ({self.tf_done}/{self.tf_total})',
                f'<code>[{tf_bar}] {tf_pct}</code>',
                f'الشموع: {self.bars_done} / {self.bars_total}',
            ]
        lines += [
            f'',
            f'W: {self.win}  L: {self.loss}  BE: {self.be}',
            f'{icon} {pnl}  WR: {wr}',
            f'',
            f'Elapsed: {self._elapsed()}',
        ]
        if self.cancelled: lines.append('<b>CANCELLED</b>')
        return '\n'.join(lines)

    async def start(self, chat_id: int) -> None:
        self.chat_id    = chat_id
        self._start_ts  = datetime.now(timezone.utc).timestamp()
        self._last_edit = self._start_ts
        payload = {'chat_id': chat_id, 'text': self._build_text(), 'parse_mode': 'HTML', 'reply_markup': self._cancel_kbd()}
        try:
            async with get_http().post(f'https://api.telegram.org/bot{TG_TOKEN}/sendMessage', json=payload) as resp:
                if resp.status == 200:
                    self.msg_id = (await resp.json())['result']['message_id']
        except Exception: pass
        self._hb_task = asyncio.create_task(self._heartbeat())

    async def _heartbeat(self) -> None:
        while not self.cancelled:
            await asyncio.sleep(self.HEARTBEAT)
            await self._edit(force=True)

    async def _edit(self, force: bool = False) -> None:
        now = datetime.now(timezone.utc).timestamp()
        if not force and (now - self._last_edit) < 3: return
        if not self.msg_id or not self.chat_id: return
        async with self._lock:
            self._last_edit = now
            payload = {'chat_id': self.chat_id, 'message_id': self.msg_id, 'text': self._build_text(), 'parse_mode': 'HTML', 'reply_markup': self._cancel_kbd() if not self.cancelled else None}
            try:
                async with get_http().post(f'https://api.telegram.org/bot{TG_TOKEN}/editMessageText', json=payload) as _: pass
            except Exception: pass

    async def set_phase(self, phase: str) -> None:
        self.phase = phase; await self._edit()

    async def set_tf(self, tf: str, bars_total: int) -> None:
        self.current_tf = tf; self.bars_done = 0
        self.bars_total = bars_total; self.phase = f'Scanning [{tf}]'
        await self._edit(force=True)

    async def tick(self, bar_n: int, win: int, loss: int, be: int, profit: float) -> None:
        self.bars_done = bar_n; self.win = win; self.loss = loss
        self.be = be; self.profit = profit; await self._edit()

    async def finish_tf(self) -> None:
        self.tf_done += 1; self.bars_done = self.bars_total
        await self._edit(force=True)

    async def done(self, final_text: str) -> None:
        if self._hb_task: self._hb_task.cancel()
        if not self.msg_id or not self.chat_id: return
        payload = {'chat_id': self.chat_id, 'message_id': self.msg_id, 'text': final_text, 'parse_mode': 'HTML'}
        try:
            async with get_http().post(f'https://api.telegram.org/bot{TG_TOKEN}/editMessageText', json=payload) as _: pass
        except Exception: pass

    async def cancel(self) -> None:
        self.cancelled = True; self.phase = 'Cancelling...'
        if self._hb_task: self._hb_task.cancel()
        await self._edit(force=True)

_bt_progress: BtProgress | None = None

def _get_cancel_kbd_for_running() -> dict:
    return {'inline_keyboard': [[{'text': 'عرض التقدم', 'callback_data': 'bt_show_progress'}], [{'text': 'إلغاء', 'callback_data': 'cancel_bt'}]]}

# ─────────────────────────────────────────────────────────────
# INDICATOR ENGINE
# ─────────────────────────────────────────────────────────────
def _ema(series: pd.Series, span: int) -> pd.Series:
    return series.ewm(span=span, adjust=False).mean()

def calculate_indicators(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty: return df
    df['ema15']  = _ema(df['close'], 15)
    df['ema50']  = _ema(df['close'], 50)
    df['ema150'] = _ema(df['close'], 150)
    
    k_period = bot_state['stoch_k']
    smooth   = bot_state['stoch_smooth']
    d_period = bot_state['stoch_d']
    
    low_min  = df['low'].rolling(k_period).min()
    high_max = df['high'].rolling(k_period).max()
    denom    = (high_max - low_min).replace(0, 1e-10)
    
    df['Fast_K'] = 100.0 * (df['close'] - low_min) / denom
    df['K'] = df['Fast_K'].rolling(window=smooth).mean()
    df['D'] = df['K'].ewm(alpha=1.0/d_period, adjust=False).mean()
    
    tr = pd.concat([
        (df['high'] - df['low']).abs(),
        (df['high'] - df['close'].shift()).abs(),
        (df['low']  - df['close'].shift()).abs(),
    ], axis=1).max(axis=1)
    df['atr'] = tr.rolling(14).mean().bfill()
    return df

# ─────────────────────────────────────────────────────────────
# SIGNAL GENERATION
# ─────────────────────────────────────────────────────────────
def get_stoch_signals(prev_k, prev_d, curr_k, curr_d) -> tuple:
    mode = bot_state['strategy_mode']
    if mode == 'STOCH_NEW':
        buy_deep  = (prev_k <= 10)      and (curr_k > 10)  and bot_state['use_stoch_deep']
        buy_mid   = (10 < prev_k <= 15) and (curr_k > 15)  and bot_state['use_stoch_mid']
        buy_shal  = (15 < prev_k <= 20) and (curr_k > 20)  and bot_state['use_stoch_shal']
        sell_deep = (prev_k >= 90)      and (curr_k < 90)  and bot_state['use_stoch_deep']
        sell_mid  = (85 <= prev_k < 90) and (curr_k < 85)  and bot_state['use_stoch_mid']
        sell_shal = (80 <= prev_k < 85) and (curr_k < 80)  and bot_state['use_stoch_shal']
    else:
        k_up   = (prev_k < prev_d) and (curr_k >= curr_d)
        k_dn   = (prev_k > prev_d) and (curr_k <= curr_d)
        avg_k  = (prev_k + curr_k) / 2.0
        buy_deep  = k_up and (avg_k <= 10)       and bot_state['use_stoch_deep']
        buy_mid   = k_up and (10 < avg_k <= 15)  and bot_state['use_stoch_mid']
        buy_shal  = k_up and (15 < avg_k <= 20)  and bot_state['use_stoch_shal']
        sell_deep = k_dn and (avg_k >= 90)       and bot_state['use_stoch_deep']
        sell_mid  = k_dn and (85 <= avg_k < 90)  and bot_state['use_stoch_mid']
        sell_shal = k_dn and (80 <= avg_k < 85)  and bot_state['use_stoch_shal']
    buy_sig  = buy_deep  or buy_mid  or buy_shal
    sell_sig = sell_deep or sell_mid or sell_shal
    b_label  = 'DEEP(10)' if buy_deep  else 'MID(15)'  if buy_mid  else 'SHAL(20)'
    s_label  = 'DEEP(90)' if sell_deep else 'MID(85)'  if sell_mid else 'SHAL(80)'
    return buy_sig, sell_sig, b_label, s_label

def compute_trend_ok(df, i, curr) -> tuple:
    mode = bot_state['filter_mode']
    if mode == 'NO_MA': return True, True
    
    # 1. الترتيب الأساسي
    b_ema = curr['ema50'] > curr['ema150']
    s_ema = curr['ema150'] > curr['ema50']
    
    # 2. فجوة الأمان (MA Gap Filter)
    gap_ok = True
    gap_val = abs(curr['ema50'] - curr['ema150'])
    
    if bot_state['use_ma_gap_fixed']:
        if gap_val < (bot_state['ma_gap_pips'] * bot_state['pip_value']):
            gap_ok = False
            
    if bot_state['use_ma_gap_atr']:
        if gap_val < (curr['atr'] * bot_state['ma_gap_atr_mult']):
            gap_ok = False
            
    b_ema = b_ema and gap_ok
    s_ema = s_ema and gap_ok
    
    if mode == 'SIMPLE': return b_ema, s_ema
    
    ma_buy  = curr['ema15'] > curr['ema50'] > curr['ema150']
    ma_sell = curr['ema15'] < curr['ema50'] < curr['ema150']
    return (b_ema and ma_buy), (s_ema and ma_sell)

def _get_signal_for_bar(df, i, curr, prev) -> tuple:
    trend_buy, trend_sell = compute_trend_ok(df, i, curr)
    raw_buy, raw_sell, b_lbl, s_lbl = get_stoch_signals(prev['K'], prev['D'], curr['K'], curr['D'])
    buy_sig  = raw_buy  and trend_buy
    sell_sig = raw_sell and trend_sell
    return buy_sig, sell_sig, (b_lbl if buy_sig else s_lbl)

# ─────────────────────────────────────────────────────────────
# DERIV WEBSOCKET FETCHER
# ─────────────────────────────────────────────────────────────
_TF_MAP = {
    '1m': 60, '2m': 120, '3m': 180, '5m': 300, '10m': 600,
    '15m': 900, '30m': 1800, '1h': 3600, '2h': 7200
}

_deriv_sem: asyncio.Semaphore | None = None

def _get_deriv_sem():
    global _deriv_sem
    if _deriv_sem is None: _deriv_sem = asyncio.Semaphore(2)
    return _deriv_sem

async def fetch_deriv_candles(granularity_str: str, count: int = 5000, end_time: datetime = None):
    resample_needed = False
    target_gran = _TF_MAP.get(granularity_str)
    
    if not target_gran:
        resample_needed = True
        fetch_gran = 60
        multiplier = int(granularity_str.replace('m', '').replace('h', '0'))
        if 'h' in granularity_str: multiplier *= 60
        fetch_count = min(count * multiplier, 15000)
    else:
        fetch_gran = target_gran
        fetch_count = min(count, 10000)

    collected = []
    current_end = int(end_time.timestamp()) if end_time else "latest"
    remaining = fetch_count

    sem = _get_deriv_sem()
    async with sem:
        try:
            async with websockets.connect(DERIV_WS_URL) as ws:
                while remaining > 0:
                    chunk = min(remaining, 5000)
                    req = {
                        "ticks_history": DERIV_SYMBOL,
                        "adjust_start_time": 1,
                        "count": chunk,
                        "end": current_end,
                        "start": 1,
                        "style": "candles",
                        "granularity": fetch_gran
                    }
                    
                    for attempt in range(3):
                        try:
                            await ws.send(json.dumps(req))
                            resp = await asyncio.wait_for(ws.recv(), timeout=10.0)
                            data = json.loads(resp)
                            
                            if 'error' in data:
                                c_log(f"Deriv API Error: {data['error']['message']}")
                                break
                            
                            candles = data.get('candles', [])
                            if not candles: break
                            
                            formatted = [{
                                'time': pd.to_datetime(c['epoch'], unit='s', utc=True),
                                'open': float(c['open']),
                                'high': float(c['high']),
                                'low': float(c['low']),
                                'close': float(c['close'])
                            } for c in candles]
                            
                            collected = formatted + collected
                            remaining -= len(candles)
                            
                            current_end = candles[0]['epoch'] - 1
                            
                            if len(candles) < chunk: remaining = 0
                            break
                        except Exception as e:
                            c_log(f"WS fetch error attempt {attempt+1}: {e}")
                            await asyncio.sleep(1)
                            
                    if remaining <= 0 or not candles: break
                    await asyncio.sleep(0.5)
        except Exception as e:
            c_log(f"WebSocket Connection Error: {e}")

    if resample_needed and collected:
        df = pd.DataFrame(collected)
        df.set_index('time', inplace=True)
        rule = granularity_str.replace('m', 'min').replace('h', 'h')
        resampled = df.resample(rule).agg({
            'open': 'first', 'high': 'max', 'low': 'min', 'close': 'last'
        }).dropna().reset_index()
        return resampled.to_dict('records')
        
    return collected

# ─────────────────────────────────────────────────────────────
# TELEGRAM HELPERS
# ─────────────────────────────────────────────────────────────
async def send_tg_msg(text: str, reply_markup: dict = None) -> None:
    if not bot_state['chat_id']: return
    payload = {'chat_id': bot_state['chat_id'], 'text': text, 'parse_mode': 'HTML'}
    if reply_markup: payload['reply_markup'] = reply_markup
    try:
        async with get_http().post(f'https://api.telegram.org/bot{TG_TOKEN}/sendMessage', json=payload) as resp:
            pass
    except Exception as e: c_log(f'TG send error: {e}')

async def edit_tg_msg(chat_id, message_id, text, reply_markup=None) -> None:
    payload = {'chat_id': chat_id, 'message_id': message_id, 'text': text, 'parse_mode': 'HTML'}
    if reply_markup: payload['reply_markup'] = reply_markup
    try:
        async with get_http().post(f'https://api.telegram.org/bot{TG_TOKEN}/editMessageText', json=payload) as resp:
            pass
    except Exception as e: c_log(f'TG edit error: {e}')

async def answer_callback(cbq_id: str, text: str = None) -> None:
    payload = {'callback_query_id': cbq_id}
    if text: payload['text'] = text
    try:
        async with get_http().post(f'https://api.telegram.org/bot{TG_TOKEN}/answerCallbackQuery', json=payload) as _: pass
    except Exception as e: c_log(f'TG callback error: {e}')

async def send_tg_document(file_path: str, caption: str) -> None:
    if not bot_state['chat_id']: return
    try:
        with open(file_path, 'rb') as f:
            data = aiohttp.FormData()
            data.add_field('chat_id', str(bot_state['chat_id']))
            data.add_field('document', f, filename=os.path.basename(file_path))
            data.add_field('caption', caption)
            async with get_http().post(f'https://api.telegram.org/bot{TG_TOKEN}/sendDocument', data=data) as resp:
                pass
    except Exception as e: c_log(f'TG doc error: {e}')

# ─────────────────────────────────────────────────────────────
# KEYBOARDS
# ─────────────────────────────────────────────────────────────
def _strat_label() -> str:
    return {'STOCH_NEW': 'STOCH-NEW', 'STOCH_OLD': 'STOCH-OLD'}[bot_state['strategy_mode']]

def _dd_status_line() -> str:
    if not bot_state['live_connected']: return 'DD: offline'
    sod = bot_state['sod_balance']
    if sod is None: return 'DD: monitoring...'
    limit_val = sod * (1 - DD_LIMIT_PCT)
    return f'DD TRIGGERED limit ${limit_val:.0f}' if bot_state['dd_triggered'] else f'DD OK  limit ${limit_val:.0f}'

def get_main_keyboard() -> dict:
    live = 'connected' if bot_state['live_connected'] else 'disconnected'
    st   = 'RUN' if bot_state['status'] == 'RUNNING' else 'PAUSE'
    bt   = 'BT running' if bot_state['is_backtesting'] else 'Backtest'
    return {'inline_keyboard': [
        [{'text': f'Server: {live}',  'callback_data': 'toggle_live_conn'}],
        [{'text': f'Bot: {st}',       'callback_data': 'toggle_status'},
         {'text': f'Strategy: {_strat_label()}', 'callback_data': 'cycle_strategy'}],
        [{'text': 'Filters',          'callback_data': 'menu_filters'},
         {'text': 'Timeframes',       'callback_data': 'menu_tfs'}],
        [{'text': 'Market Report',    'callback_data': 'report'},
         {'text': 'Account',          'callback_data': 'account'}],
        [{'text': 'Risk Settings',    'callback_data': 'menu_settings'},
         {'text': bt,                 'callback_data': 'menu_backtest'}],
        [{'text': _dd_status_line(),  'callback_data': 'dd_status'}],
        [{'text': 'Close All Trades', 'callback_data': 'close_all'}],
    ]}

def get_filters_keyboard() -> dict:
    fm = bot_state['filter_mode']
    fi = {k: '[X]' if fm == k else '[ ]' for k in ('FULL', 'SIMPLE', 'NO_MA')}
    dp = 'ON' if bot_state['use_stoch_deep'] else 'OFF'
    md = 'ON' if bot_state['use_stoch_mid']  else 'OFF'
    sh = 'ON' if bot_state['use_stoch_shal'] else 'OFF'
    d_i = 'ON' if bot_state['use_danger_filter'] else 'OFF'
    k, s, d = bot_state['stoch_k'], bot_state['stoch_smooth'], bot_state['stoch_d']
    return {'inline_keyboard': [
        [{'text': '-- Trend Filter --', 'callback_data': 'noop'}],
        [{'text': f"{fi['FULL']} FULL: ema15+ema50+ema150", 'callback_data': 'set_filter_full'}],
        [{'text': f"{fi['SIMPLE']} SIMPLE: ema50+ema150",   'callback_data': 'set_filter_simple'}],
        [{'text': f"{fi['NO_MA']} NO MA",                   'callback_data': 'set_filter_noma'}],
        [{'text': '-- MA Gap Filter --', 'callback_data': 'noop'}],
        [{'text': '⚙️ إعدادات فجوة الموفينجات (MA Gap)', 'callback_data': 'menu_ma_gap'}],
        [{'text': '-- Stochastic Levels --', 'callback_data': 'noop'}],
        [{'text': f'Stoch({k},{s},{d})', 'callback_data': 'menu_stoch_settings'}],
        [{'text': f'DEEP 10/90: {dp}',   'callback_data': 'toggle_stoch_deep'},
         {'text': f'MID  15/85: {md}',   'callback_data': 'toggle_stoch_mid'},
         {'text': f'SHAL 20/80: {sh}',   'callback_data': 'toggle_stoch_shal'}],
        [{'text': '-- Time Filter (Danger Zones) --', 'callback_data': 'noop'}],
        [{'text': f'Block (13, 18, 21, 22 Damascus): {d_i}', 'callback_data': 'toggle_danger'}],
        [{'text': 'Back', 'callback_data': 'menu_main'}],
    ]}

def get_ma_gap_keyboard() -> dict:
    fx_i = 'ON' if bot_state['use_ma_gap_fixed'] else 'OFF'
    at_i = 'ON' if bot_state['use_ma_gap_atr'] else 'OFF'
    pips = bot_state['ma_gap_pips']
    mult = bot_state['ma_gap_atr_mult']
    return {'inline_keyboard': [
        [{'text': f'Fixed Pips Gap: {fx_i}', 'callback_data': 'toggle_ma_gap_fixed'}],
        [{'text': '-', 'callback_data': 'dec_gap_pips'},
         {'text': f'{pips} Pips', 'callback_data': 'noop'},
         {'text': '+', 'callback_data': 'inc_gap_pips'}],
        [{'text': f'ATR Multiplier Gap: {at_i}', 'callback_data': 'toggle_ma_gap_atr'}],
        [{'text': '-', 'callback_data': 'dec_gap_atr'},
         {'text': f'{mult}x ATR', 'callback_data': 'noop'},
         {'text': '+', 'callback_data': 'inc_gap_atr'}],
        [{'text': 'Back', 'callback_data': 'menu_filters'}],
    ]}

def get_stoch_settings_keyboard() -> dict:
    k = bot_state['stoch_k']; s = bot_state['stoch_smooth']; d = bot_state['stoch_d']
    return {'inline_keyboard': [
        [{'text': f'Stoch({k},{s},{d})', 'callback_data': 'noop'}],
        [{'text': '/stoch K S D  e.g. /stoch 14 3 3', 'callback_data': 'noop'}],
        [{'text': '-- K Period --', 'callback_data': 'noop'}],
        [{'text': '-', 'callback_data': 'dec_stoch_k'}, {'text': f'K={k}', 'callback_data': 'noop'}, {'text': '+', 'callback_data': 'inc_stoch_k'}],
        [{'text': '-- Smooth --', 'callback_data': 'noop'}],
        [{'text': '-', 'callback_data': 'dec_stoch_s'}, {'text': f'S={s}', 'callback_data': 'noop'}, {'text': '+', 'callback_data': 'inc_stoch_s'}],
        [{'text': '-- D Period --', 'callback_data': 'noop'}],
        [{'text': '-', 'callback_data': 'dec_stoch_d'}, {'text': f'D={d}', 'callback_data': 'noop'}, {'text': '+', 'callback_data': 'inc_stoch_d'}],
        [{'text': '5,5,5', 'callback_data': 'preset_5_5_5'}, {'text': '14,3,3', 'callback_data': 'preset_14_3_3'}, {'text': '10,3,3', 'callback_data': 'preset_10_3_3'}],
        [{'text': 'Back', 'callback_data': 'menu_filters'}],
    ]}

def get_tf_keyboard() -> dict:
    rows, row = [], []
    for tf in bot_state['timeframes']:
        icon = 'ON' if bot_state['active_tfs'][tf] else 'OFF'
        row.append({'text': f'{tf}: {icon}', 'callback_data': f'toggle_tf_{tf}'})
        if len(row) == 3: rows.append(row); row = []
    if row: rows.append(row)
    rows.append([{'text': 'Back', 'callback_data': 'menu_main'}])
    return {'inline_keyboard': rows}

def get_settings_keyboard() -> dict:
    be_i  = 'ON' if bot_state['use_be']         else 'OFF'
    atr_i = 'ON' if bot_state['use_atr']        else 'OFF'
    spr_i = 'ON' if bot_state['use_max_spread'] else 'OFF'
    return {'inline_keyboard': [
        [{'text': f'BE 20p: {be_i}',               'callback_data': 'toggle_be'}],
        [{'text': f'ATR targets: {atr_i}',          'callback_data': 'toggle_atr'}],
        [{'text': f'Spread guard {bot_state["max_spread_pips"]}p: {spr_i}', 'callback_data': 'toggle_spread'}],
        [{'text': f'Lot: {bot_state["lot_size"]:.2f}', 'callback_data': 'noop'}],
        [{'text': '+Lot', 'callback_data': 'inc_lot'}, {'text': '-Lot', 'callback_data': 'dec_lot'}],
        [{'text': 'Edit TP/SL per TF', 'callback_data': 'view_tpsl'}],
        [{'text': 'Back',              'callback_data': 'menu_main'}],
    ]}

def get_tpsl_overview_keyboard() -> dict:
    rows = [[{'text': '-- Tap a TF to edit --', 'callback_data': 'noop'}]]
    for tf in bot_state['timeframes']:
        icon = 'ON' if bot_state['active_tfs'][tf] else 'OFF'
        tp = bot_state['tp_pips'][tf]; sl = bot_state['sl_pips'][tf]
        rows.append([
            {'text': f'[{icon}] [{tf}] TP:{tp} SL:{sl}', 'callback_data': 'noop'},
            {'text': 'Edit', 'callback_data': f'tpsl_edit_{tf}'},
        ])
    rows.append([{'text': '/set 1m sl 75', 'callback_data': 'noop'}])
    rows.append([{'text': 'Back', 'callback_data': 'menu_settings'}])
    return {'inline_keyboard': rows}

def get_tpsl_edit_keyboard(tf: str) -> dict:
    tp = bot_state['tp_pips'][tf]; sl = bot_state['sl_pips'][tf]
    rr = round(tp / sl, 2) if sl else 0
    return {'inline_keyboard': [
        [{'text': f'[{tf}] TP:{tp}p SL:{sl}p RR:1:{rr}', 'callback_data': 'noop'}],
        [{'text': 'Take Profit', 'callback_data': 'noop'}],
        [{'text': f'-10({tp-10})', 'callback_data': f'dec_tp10_{tf}'},
         {'text': f'TP={tp}',     'callback_data': 'noop'},
         {'text': f'+10({tp+10})', 'callback_data': f'inc_tp10_{tf}'}],
        [{'text': f'-5({tp-5})',  'callback_data': f'dec_tp5_{tf}'},
         {'text': '-',            'callback_data': 'noop'},
         {'text': f'+5({tp+5})',  'callback_data': f'inc_tp5_{tf}'}],
        [{'text': 'Stop Loss', 'callback_data': 'noop'}],
        [{'text': f'-10({sl-10})', 'callback_data': f'dec_sl10_{tf}'},
         {'text': f'SL={sl}',     'callback_data': 'noop'},
         {'text': f'+10({sl+10})', 'callback_data': f'inc_sl10_{tf}'}],
        [{'text': f'-5({sl-5})',  'callback_data': f'dec_sl5_{tf}'},
         {'text': '-',            'callback_data': 'noop'},
         {'text': f'+5({sl+5})',  'callback_data': f'inc_sl5_{tf}'}],
        [{'text': f'/set {tf} tp|sl value', 'callback_data': 'noop'}],
        [{'text': 'Back', 'callback_data': 'view_tpsl'}],
    ]}

def get_backtest_keyboard() -> dict:
    if bot_state['is_backtesting']:
        return {'inline_keyboard': [
            [{'text': 'BT running...', 'callback_data': 'bt_show_progress'}],
            [{'text': 'Cancel',        'callback_data': 'cancel_bt'}],
            [{'text': 'Back',          'callback_data': 'menu_main'}],
        ]}
    return {'inline_keyboard': [
        [{'text': '1 day',  'callback_data': 'bto_1'},
         {'text': '3 days', 'callback_data': 'bto_3'},
         {'text': '7 days', 'callback_data': 'bto_7'}],
        [{'text': 'Advanced 7 days', 'callback_data': 'bto_adv_7'}],
        [{'text': 'Back', 'callback_data': 'menu_main'}],
    ]}

# ─────────────────────────────────────────────────────────────
# BACKTEST SHARED HELPERS
# ─────────────────────────────────────────────────────────────
def _build_trade_row(tf, is_buy, label, entry_t, exit_t, act_ent, tp_p, sl_p, outcome, p_usd):
    pv = bot_state['pip_value']
    # تم إزالة +3 ساعات لتتطابق الأوقات مع شاشة Deriv (توقيت السيرفر)
    return {
        'Timeframe':   tf,
        'Type':        ('BUY' if is_buy else 'SELL') + f' [{label}]',
        'Entry Time':  entry_t.strftime('%Y-%m-%d %H:%M'),
        'Exit Time':   exit_t.strftime('%Y-%m-%d %H:%M'),
        'Entry Price': round(act_ent, 2),
        'TP': tp_p, 'SL': sl_p,
        'Pips': (round(abs(act_ent - (tp_p if outcome == 'WIN' else sl_p)) / pv, 1)
                 if outcome in ('WIN', 'LOSS') else 0),
        'Outcome': outcome, 'Profit ($)': p_usd,
    }

def _simulate_trade_in_memory(is_buy, act_ent, tp_p, sl_p, eff_tp, entry_t, minute_candles):
    pv      = bot_state['pip_value']
    be_act  = False
    be_tgt  = act_ent + (1 if is_buy else -1) * 20 * pv
    max_ext = entry_t + timedelta(hours=72)
    outcome = 'EXPIRED'; exit_t = max_ext
    for vc in minute_candles:
        t = vc['time']
        if t < entry_t: continue
        if t > max_ext: break
        if is_buy:
            if bot_state['use_be'] and not be_act and vc['high'] >= be_tgt:
                sl_p = act_ent; be_act = True
            if vc['low'] <= sl_p:
                outcome = 'BREAK-EVEN' if be_act else 'LOSS'; exit_t = t; break
            if vc['high'] >= eff_tp:
                outcome = 'WIN'; exit_t = t; break
        else:
            if bot_state['use_be'] and not be_act and vc['low'] <= be_tgt:
                sl_p = act_ent; be_act = True
            if vc['high'] >= sl_p:
                outcome = 'BREAK-EVEN' if be_act else 'LOSS'; exit_t = t; break
            if vc['low'] <= eff_tp:
                outcome = 'WIN'; exit_t = t; break
    return outcome, exit_t, sl_p

def _calc_pnl(outcome, act_ent, tp_p, sl_p):
    if outcome == 'BREAK-EVEN': return 0.0
    if outcome in ('WIN', 'LOSS'):
        exit_p = tp_p if outcome == 'WIN' else sl_p
        raw    = abs(act_ent - exit_p) * 100 * bot_state['lot_size']
        return round(raw, 2) * (1 if outcome == 'WIN' else -1)
    return 0.0

def _entry_params(curr, is_buy, tf):
    m       = 1 if is_buy else -1
    act_ent = curr['open'] + m * bot_state['spread_pips'] * bot_state['pip_value']
    tp_dist = (curr['atr'] * bot_state['atr_mult_tp'] if bot_state['use_atr']
               else bot_state['tp_pips'][tf] * bot_state['pip_value'])
    sl_dist = (curr['atr'] * bot_state['atr_mult_sl'] if bot_state['use_atr']
               else bot_state['sl_pips'][tf] * bot_state['pip_value'])
    tp_p   = round(act_ent + m * tp_dist, 2)
    sl_p   = round(act_ent - m * sl_dist, 2)
    tol    = bot_state['tp_tolerance_pips'] * bot_state['pip_value']
    eff_tp = (tp_p - tol) if is_buy else (tp_p + tol)
    return act_ent, tp_p, sl_p, eff_tp

def _build_final_summary(desc, trade_logs, blocked_logs, win_count, loss_count,
                          be_count, total_win, total_loss, total_prof, peak_equity, max_dd):
    total_trades = win_count + loss_count
    win_rate     = round(win_count / total_trades * 100, 1) if total_trades else 0
    dd_pct       = round(max_dd / peak_equity * 100, 1) if peak_equity else 0
    icon         = 'PROFIT' if total_prof >= 0 else 'LOSS'
    return (
        f'<b>Backtest Complete</b>\n{desc}\n'
        f'Net: {icon} ${round(total_prof, 2)}\n'
        f'Win: +${round(total_win, 2)} ({win_count})\n'
        f'Loss: -${abs(round(total_loss, 2))} ({loss_count})\n'
        f'BE: {be_count}\n'
        f'WR: {win_rate}% ({total_trades} trades)\n'
        f'Max DD: ${round(max_dd, 2)} ({dd_pct}%)\n'
        f'Sending Excel file...'
    )

async def _fetch_1m_candles_for_bt(start_dt, tf_end):
    total_min = int((tf_end - start_dt).total_seconds() / 60) + 72 * 60 + 60
    m1_raw = await fetch_deriv_candles('1m', count=total_min, end_time=tf_end)
    return sorted(m1_raw, key=lambda x: x['time'])

# ─────────────────────────────────────────────────────────────
# BACKTESTING
# ─────────────────────────────────────────────────────────────
async def run_deriv_backtest(start_dt: datetime) -> None:
    global _bt_progress
    if bot_state['is_backtesting']: return
    bot_state['is_backtesting'] = True
    active_tfs = [tf for tf in bot_state['timeframes'] if bot_state['active_tfs'][tf]]
    desc       = f"{bot_state['strategy_mode']} / {bot_state['filter_mode']}"
    fname      = f"BT_{datetime.now().strftime('%H%M%S')}.xlsx"
    prog = BtProgress(label=desc, active_tfs=active_tfs)
    _bt_progress = prog
    await prog.start(bot_state['chat_id'])
    trade_logs, blocked_logs = [], []
    total_prof = peak_equity = max_dd = 0.0
    total_win  = total_loss  = 0.0
    win_count  = loss_count  = be_count = 0
    try:
        for tf in active_tfs:
            if prog.cancelled: break
            await asyncio.sleep(0)
            await prog.set_phase(f'إحماء الذاكرة وجلب الشموع [{tf}]...')
            
            c_data = await fetch_deriv_candles(tf, count=10000)
            if len(c_data) < 150: await prog.finish_tf(); continue
            df         = calculate_indicators(pd.DataFrame(c_data).sort_values('time').reset_index(drop=True))
            safe_start = max(10, 3) # تم إزالة الاعتماد على cons_count هنا لعدم وجوده
            
            await prog.set_phase(f'Loading 1m for [{tf}]...')
            minute_candles = await _fetch_1m_candles_for_bt(start_dt, datetime.now(timezone.utc))
            bar_index = df[df['time'] >= start_dt].index.tolist()
            await prog.set_tf(tf, len(bar_index))
            for loop_pos, i in enumerate(bar_index):
                if prog.cancelled: break
                if i < safe_start: continue
                if loop_pos % 50 == 0:
                    await asyncio.sleep(0)
                    await prog.tick(loop_pos, win_count, loss_count, be_count, total_prof)
                curr = df.loc[i]; prev = df.loc[i - 1]
                if bot_state['use_danger_filter'] and is_blocked_time(curr['time']): continue
                buy_sig, sell_sig, label = _get_signal_for_bar(df, i, curr, prev)
                raw_buy, raw_sell, b_lbl, s_lbl = get_stoch_signals(prev['K'], prev['D'], curr['K'], curr['D'])
                ts = curr['time'].strftime('%Y-%m-%d %H:%M') # توقيت السيرفر
                if not buy_sig  and raw_buy:
                    blocked_logs.append({'Timeframe': tf, 'Type': f'BUY BLOCKED ({b_lbl})', 'Entry Time': ts, 'Entry Price': curr['close'], 'Reason': f'REJECTED ({bot_state["filter_mode"]} + MA GAP)'})
                if not sell_sig and raw_sell:
                    blocked_logs.append({'Timeframe': tf, 'Type': f'SELL BLOCKED ({s_lbl})', 'Entry Time': ts, 'Entry Price': curr['close'], 'Reason': f'REJECTED ({bot_state["filter_mode"]} + MA GAP)'})
                if not (buy_sig or sell_sig) or i + 1 >= len(df): continue
                next_c = df.loc[i + 1]; entry_t = next_c['time']; is_buy = bool(buy_sig)
                sig_bar = next_c.copy(); sig_bar['atr'] = curr['atr']
                act_ent, tp_p, sl_p, eff_tp = _entry_params(sig_bar, is_buy, tf)
                outcome, exit_t, sl_p = _simulate_trade_in_memory(is_buy, act_ent, tp_p, sl_p, eff_tp, entry_t, minute_candles)
                p_usd = _calc_pnl(outcome, act_ent, tp_p, sl_p)
                if outcome == 'BREAK-EVEN': be_count += 1
                elif outcome == 'WIN':      total_win += p_usd; win_count += 1
                elif outcome == 'LOSS':     total_loss += p_usd; loss_count += 1
                total_prof += p_usd; peak_equity = max(peak_equity, total_prof)
                max_dd = max(max_dd, peak_equity - total_prof)
                trade_logs.append(_build_trade_row(tf, is_buy, label, entry_t, exit_t, act_ent, tp_p, sl_p, outcome, p_usd))
            await prog.finish_tf()
            await prog.tick(len(bar_index), win_count, loss_count, be_count, total_prof)
        if not trade_logs:
            await prog.done('No trades found.'); return
        await prog.done(_build_final_summary(desc, trade_logs, blocked_logs, win_count, loss_count,
                                              be_count, total_win, total_loss, total_prof, peak_equity, max_dd))
        total_trades = win_count + loss_count
        win_rate  = round(win_count / total_trades * 100, 1) if total_trades else 0
        dd_pct    = round(max_dd / peak_equity * 100, 1) if peak_equity else 0
        summary   = {
            'Item': ['Win', 'Loss', 'Net', 'WinRate', 'MaxDD', 'BE', 'Strategy'],
            'Value': [f'{win_count} +${round(total_win,2)}', f'{loss_count} -${abs(round(total_loss,2))}',
                      f'${round(total_prof,2)}', f'{win_rate}% ({total_trades})',
                      f'${round(max_dd,2)} ({dd_pct}%)', str(be_count), desc],
        }
        with pd.ExcelWriter(fname, engine='openpyxl') as writer:
            pd.DataFrame(trade_logs).to_excel(writer, sheet_name='Trades', index=False)
            pd.DataFrame(summary).to_excel(writer, sheet_name='Summary', index=False)
            if blocked_logs: pd.DataFrame(blocked_logs).to_excel(writer, sheet_name='Blocked', index=False)
            _style_sheet(writer.sheets['Trades'])
        await send_tg_document(fname, f'Backtest | {desc} | Net:${round(total_prof,2)} WR:{win_rate}%')
        os.remove(fname)
    except Exception as e:
        await prog.done(f'ERROR: {e}'); c_log(f'Backtest error: {e}')
    finally:
        bot_state['is_backtesting'] = False; _bt_progress = None

async def run_advanced_backtest(days: int = 7) -> None:
    global _bt_progress
    if bot_state['is_backtesting']: return
    bot_state['is_backtesting'] = True
    start_dt   = datetime.now(timezone.utc) - timedelta(days=days)
    active_tfs = [tf for tf in bot_state['timeframes'] if bot_state['active_tfs'][tf]]
    desc       = f"{bot_state['strategy_mode']} / {bot_state['filter_mode']}"
    prog = BtProgress(label=f'{desc} ({days}d)', active_tfs=active_tfs, is_advanced=True)
    _bt_progress = prog
    await prog.start(bot_state['chat_id'])
    trade_logs, blocked_logs = [], []
    total_prof  = peak_equity = max_dd = 0.0
    total_win   = total_loss  = 0.0
    win_count   = loss_count  = be_count = 0
    long_win    = long_loss   = short_win = short_loss = 0
    all_profits = []
    consec_win  = consec_loss = max_cw = max_cl = 0
    max_cw_usd  = max_cl_usd = cur_w = cur_l = 0.0
    try:
        for tf in active_tfs:
            if prog.cancelled: break
            await asyncio.sleep(0)
            await prog.set_phase(f'إحماء الذاكرة وجلب الشموع [{tf}]...')
            c_data = await fetch_deriv_candles(tf, count=10000)
            if len(c_data) < 150: await prog.finish_tf(); continue
            df         = calculate_indicators(pd.DataFrame(c_data).sort_values('time').reset_index(drop=True))
            safe_start = max(10, 3)
            await prog.set_phase(f'Loading 1m for [{tf}]...')
            minute_candles = await _fetch_1m_candles_for_bt(start_dt, datetime.now(timezone.utc))
            bar_index = df[df['time'] >= start_dt].index.tolist()
            await prog.set_tf(tf, len(bar_index))
            for loop_pos, i in enumerate(bar_index):
                if prog.cancelled: break
                if i < safe_start: continue
                if loop_pos % 50 == 0:
                    await asyncio.sleep(0)
                    await prog.tick(loop_pos, win_count, loss_count, be_count, total_prof)
                curr = df.loc[i]; prev = df.loc[i - 1]
                if bot_state['use_danger_filter'] and is_blocked_time(curr['time']): continue
                buy_sig, sell_sig, label = _get_signal_for_bar(df, i, curr, prev)
                raw_buy, raw_sell, b_lbl, s_lbl = get_stoch_signals(prev['K'], prev['D'], curr['K'], curr['D'])
                ts = curr['time'].strftime('%Y-%m-%d %H:%M')
                if not buy_sig  and raw_buy:
                    blocked_logs.append({'Timeframe': tf, 'Type': f'BUY BLOCKED ({b_lbl})', 'Entry Time': ts, 'Entry Price': curr['close'], 'Reason': f'REJECTED ({bot_state["filter_mode"]} + MA GAP)'})
                if not sell_sig and raw_sell:
                    blocked_logs.append({'Timeframe': tf, 'Type': f'SELL BLOCKED ({s_lbl})', 'Entry Time': ts, 'Entry Price': curr['close'], 'Reason': f'REJECTED ({bot_state["filter_mode"]} + MA GAP)'})
                if not (buy_sig or sell_sig) or i + 1 >= len(df): continue
                next_c = df.loc[i + 1]; entry_t = next_c['time']; is_buy = bool(buy_sig)
                sig_bar = next_c.copy(); sig_bar['atr'] = curr['atr']
                act_ent, tp_p, sl_p, eff_tp = _entry_params(sig_bar, is_buy, tf)
                outcome, exit_t, sl_p = _simulate_trade_in_memory(is_buy, act_ent, tp_p, sl_p, eff_tp, entry_t, minute_candles)
                p_usd = _calc_pnl(outcome, act_ent, tp_p, sl_p)
                if outcome == 'WIN':
                    total_win += p_usd; win_count += 1
                    consec_win += 1; cur_w += p_usd; consec_loss = 0; cur_l = 0.0
                    if consec_win > max_cw: max_cw = consec_win; max_cw_usd = cur_w
                    if is_buy: long_win  += 1
                    else:      short_win += 1
                elif outcome == 'LOSS':
                    total_loss  += p_usd; loss_count  += 1
                    consec_loss += 1; cur_l += p_usd; consec_win = 0; cur_w = 0.0
                    if consec_loss > max_cl: max_cl = consec_loss; max_cl_usd = cur_l
                    if is_buy: long_loss  += 1
                    else:      short_loss += 1
                elif outcome == 'BREAK-EVEN':
                    be_count += 1
                total_prof += p_usd; peak_equity = max(peak_equity, total_prof)
                max_dd = max(max_dd, peak_equity - total_prof); all_profits.append(p_usd)
                row = _build_trade_row(tf, is_buy, label, entry_t, exit_t, act_ent, tp_p, sl_p, outcome, p_usd)
                row['Hour_Damascus'] = (curr['time'].hour + 3) % 24
                row['Weekday']       = curr['time'].strftime('%a')
                trade_logs.append(row)
            await prog.finish_tf()
            await prog.tick(len(bar_index), win_count, loss_count, be_count, total_prof)
        if not trade_logs:
            await prog.done('No trades found.'); return
        total_trades    = win_count + loss_count
        win_rate        = round(win_count / total_trades * 100, 1) if total_trades else 0
        dd_pct          = round(max_dd / peak_equity * 100, 1)     if peak_equity  else 0
        profit_factor   = round(total_win / abs(total_loss), 2)    if total_loss   else 999.0
        expected_payoff = round(total_prof / total_trades, 2)       if total_trades else 0
        recovery_factor = round(total_prof / max_dd, 2)             if max_dd       else 999.0
        wins_only       = [p for p in all_profits if p > 0]
        losses_only     = [p for p in all_profits if p < 0]
        avg_win         = round(sum(wins_only)   / len(wins_only),   2) if wins_only   else 0
        avg_loss        = round(sum(losses_only) / len(losses_only), 2) if losses_only else 0
        largest_win     = round(max(wins_only),  2) if wins_only   else 0
        largest_loss    = round(min(losses_only), 2) if losses_only else 0
        df_t        = pd.DataFrame(trade_logs)
        actv        = df_t[df_t['Outcome'].isin(['WIN', 'LOSS'])]
        hour_counts = actv.groupby('Hour_Damascus').size()
        day_counts  = actv.groupby('Weekday').size()
        def barchart(dd, width=18):
            if not dd: return '(no data)'
            mx = max(dd.values())
            return '\n'.join(f'  {str(k):>4} |{"#" * int(v/mx*width):<{width}}| {v}' for k, v in sorted(dd.items()))
        await prog.done(_build_final_summary(desc, trade_logs, blocked_logs, win_count, loss_count,
                                              be_count, total_win, total_loss, total_prof, peak_equity, max_dd))
        report = (
            f'<b>Advanced Report {days}d</b>\n{desc}\n'
            f'Net: ${round(total_prof,2)} | PF:{profit_factor} | RF:{recovery_factor}\n'
            f'DD: ${round(max_dd,2)} ({dd_pct}%)\n'
            f'{total_trades} trades | W:{win_count}({win_rate}%) | L:{loss_count}\n'
            f'Long W/L:{long_win}/{long_loss} | Short W/L:{short_win}/{short_loss}\n'
            f'BE:{be_count}\n'
            f'MaxWin:+${largest_win} MaxLoss:${largest_loss}\n'
            f'AvgWin:+${avg_win} AvgLoss:${avg_loss}\n'
            f'Streak W:{max_cw}(+${round(max_cw_usd,2)}) L:{max_cl}(-${abs(round(max_cl_usd,2))})\n'
            f'<b>By Hour:</b>\n<pre>{barchart(hour_counts.to_dict())}</pre>\n'
            f'<b>By Day:</b>\n<pre>{barchart(day_counts.to_dict())}</pre>'
        )
        await send_tg_msg(report)
        xlsx_adv = f"ADV_{datetime.now().strftime('%H%M%S')}.xlsx"
        df_exec  = df_t.drop(columns=['Hour_Damascus', 'Weekday'], errors='ignore')
        stats = {
            'Metric': ['Net','Win','Loss','PF','EP','RF','MaxDD','DD%','Trades','Wins','Losses','WR','BE','LongWL','ShortWL','MaxWin','MaxLoss','AvgWin','AvgLoss','WinStreak','LossStreak','Strategy'],
            'Value':  [f'${round(total_prof,2)}',f'+${round(total_win,2)}',f'-${abs(round(total_loss,2))}',profit_factor,expected_payoff,recovery_factor,f'${round(max_dd,2)}',f'{dd_pct}%',total_trades,win_count,loss_count,f'{win_rate}%',be_count,f'{long_win}/{long_loss}',f'{short_win}/{short_loss}',f'+${largest_win}',f'${largest_loss}',f'+${avg_win}',f'${avg_loss}',f'{max_cw}(+${round(max_cw_usd,2)})',f'{max_cl}(-${abs(round(max_cl_usd,2))})',desc],
        }
        with pd.ExcelWriter(xlsx_adv, engine='openpyxl') as writer:
            df_exec.to_excel(writer, sheet_name='Trades', index=False)
            pd.DataFrame(stats).to_excel(writer, sheet_name='Stats', index=False)
            if blocked_logs: pd.DataFrame(blocked_logs).to_excel(writer, sheet_name='Blocked', index=False)
            _style_sheet(writer.sheets['Trades'])
        await send_tg_document(xlsx_adv, f'Advanced Report {days}d | {desc}')
        os.remove(xlsx_adv)
    except Exception as e:
        await prog.done(f'ERROR: {e}'); c_log(f'Advanced BT error: {e}')
    finally:
        bot_state['is_backtesting'] = False; _bt_progress = None

def _style_sheet(ws) -> None:
    from openpyxl.styles import PatternFill, Font
    green  = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    red    = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    header = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
    for cell in ws[1]: cell.fill = header; cell.font = Font(color='FFFFFF', bold=True)
    outcome_col = next((i + 1 for i, c in enumerate(ws[1]) if c.value == 'Outcome'), 9)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        val = str(row[outcome_col - 1].value) if len(row) >= outcome_col else ''
        if val == 'WIN':
            for cell in row: cell.fill = green
        elif val == 'LOSS':
            for cell in row: cell.fill = red
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 28)

# ─────────────────────────────────────────────────────────────
# LIVE MONITORS
# ─────────────────────────────────────────────────────────────
async def position_monitor() -> None:
    while True:
        try:
            if bot_state['live_connected'] and bot_state['use_be'] and bot_state['connection_obj']:
                pv        = bot_state['pip_value']
                positions = await bot_state['connection_obj'].get_positions()
                for p in positions:
                    if p['symbol'] != bot_state['symbol']: continue
                    op, tp, sl, cp = p['openPrice'], p.get('takeProfit'), p.get('stopLoss'), p['currentPrice']
                    if tp and sl != op and abs(cp - op) >= 20 * pv:
                        is_buy = tp > op
                        if (is_buy and cp > op) or (not is_buy and cp < op):
                            await bot_state['connection_obj'].modify_position(p['id'], stop_loss=op)
                            await send_tg_msg(f"BE activated — trade {p['id']}")
        except Exception as e:
            c_log(f'Position monitor error: {e}')
        await asyncio.sleep(5)

async def timeframe_scanner(tf: str) -> None:
    c_log(f'Scanner [{tf}] started.')
    while True:
        try:
            if not (bot_state['status'] == 'RUNNING' and bot_state['active_tfs'][tf]):
                await asyncio.sleep(10); continue

            if not (bot_state['live_connected'] and bot_state['account_obj']):
                bot_state['market_data'][tf] = 'Offline'
                await asyncio.sleep(5); continue

            if bot_state['dd_triggered']:
                bot_state['market_data'][tf] = 'DD PAUSED'
                await asyncio.sleep(30); continue

            if bot_state['cache'][tf].empty:
                await send_tg_msg(
                    f"⏳ <b>[نظام الذاكرة Deriv]</b>\n"
                    f"بدء إحماء فريم <b>[{tf}]</b>...\n"
                    f"جاري سحب 5000 شمعة تاريخية لضبط معيار MT5."
                )
                c_log(f'[{tf}] WARM-UP: Fetching 5000 historical candles from Deriv...')
                try:
                    raw = await fetch_deriv_candles(tf, count=5000)
                    if not raw:
                        c_log(f'[{tf}] WARM-UP FAILED: No data received.')
                        await asyncio.sleep(15); continue
                        
                    bot_state['cache'][tf] = pd.DataFrame(raw)
                    await send_tg_msg(
                        f"✅ <b>[نظام الذاكرة Deriv]</b>\n"
                        f"اكتمل إحماء <b>[{tf}]</b> بنجاح!\n"
                        f"تم تخزين {len(raw)} شمعة. البوت يعمل الآن بدقة 100%."
                    )
                    c_log(f'[{tf}] WARM-UP SUCCESS: Cached {len(raw)} candles.')
                except Exception as e:
                    c_log(f'[{tf}] WARM-UP ERROR: {e}')
                    await send_tg_msg(f"❌ <b>[نظام الذاكرة Deriv]</b>\nخطأ أثناء إحماء <b>[{tf}]</b>:\n{e}\nسيعاود المحاولة...")
                    await asyncio.sleep(15); continue

            try:
                raw_new = await fetch_deriv_candles(tf, count=10)
                if not raw_new:
                    await asyncio.sleep(10); continue
                    
                df_new = pd.DataFrame(raw_new)
                df_cached = bot_state['cache'][tf]
                
                df_combined = pd.concat([df_cached, df_new], ignore_index=True)
                df_combined.drop_duplicates(subset=['time'], keep='last', inplace=True)
                
                df_combined = df_combined.tail(5000).reset_index(drop=True)
                bot_state['cache'][tf] = df_combined
                
                c_log(f'[{tf}] Pulse Update (Deriv): Appended latest candles. Total cache: {len(df_combined)}.')
                df = calculate_indicators(df_combined.copy())
            except Exception as e:
                c_log(f'[{tf}] PULSE ERROR: {e}')
                await asyncio.sleep(10); continue

            if df.empty or len(df) < 3: await asyncio.sleep(15); continue
                
            curr    = df.iloc[-2]; prev = df.iloc[-3]
            now_utc = datetime.now(timezone.utc)

            bot_state['market_data'][tf] = (
                f"{df.iloc[-1]['close']:.2f} | K:{curr['K']:.1f} D:{curr['D']:.1f}"
            )

            if bot_state['use_danger_filter'] and is_blocked_time(now_utc):
                lbl = blocked_time_label(now_utc)
                bot_state['market_data'][tf] = f'BLOCKED {lbl} | {df.iloc[-1]["close"]:.2f}'
                await asyncio.sleep(10); continue

            if bot_state['last_signal_time'][tf] == curr['time']:
                await asyncio.sleep(10); continue

            trend_buy, trend_sell = compute_trend_ok_live(df, curr)
            raw_buy, raw_sell, b_lbl, s_lbl = get_stoch_signals(prev['K'], prev['D'], curr['K'], curr['D'])
            buy_sig  = raw_buy  and trend_buy
            sell_sig = raw_sell and trend_sell
            label    = b_lbl if buy_sig else s_lbl

            if bot_state['use_max_spread'] and (buy_sig or sell_sig):
                try:
                    tick        = await bot_state['connection_obj'].get_tick(bot_state['symbol'])
                    spread_pips = (tick['ask'] - tick['bid']) / bot_state['pip_value']
                    if spread_pips > bot_state['max_spread_pips']:
                        c_log(f'[{tf}] spread {spread_pips:.1f}p > max, skip')
                        buy_sig = sell_sig = False
                except Exception: pass

            if not (buy_sig or sell_sig):
                await asyncio.sleep(10); continue

            bot_state['last_signal_time'][tf] = curr['time']
            price = df.iloc[-1]['close']
            m     = 1 if buy_sig else -1
            t_str = 'BUY' if buy_sig else 'SELL'

            tp_dist = (curr['atr'] * bot_state['atr_mult_tp'] if bot_state['use_atr']
                       else bot_state['tp_pips'][tf] * bot_state['pip_value'])
            sl_dist = (curr['atr'] * bot_state['atr_mult_sl'] if bot_state['use_atr']
                       else bot_state['sl_pips'][tf] * bot_state['pip_value'])
            tp = round(price + m * tp_dist, 2)
            sl = round(price - m * sl_dist, 2)

            try:
                if buy_sig:
                    await bot_state['connection_obj'].create_market_buy_order(
                        bot_state['symbol'], bot_state['lot_size'], stop_loss=sl, take_profit=tp)
                else:
                    await bot_state['connection_obj'].create_market_sell_order(
                        bot_state['symbol'], bot_state['lot_size'], stop_loss=sl, take_profit=tp)
                await send_tg_msg(
                    f'<b>Trade Opened</b>\n{t_str} [{tf}]\n'
                    f'Price:{price:.2f} TP:{tp} SL:{sl}\n[{label}]'
                )
            except Exception as e:
                await send_tg_msg(f'<b>Order Failed [{tf}]:</b>\n{e}')

        except Exception as e:
            c_log(f'Scanner [{tf}] error: {e}')
        await asyncio.sleep(10)

# ─────────────────────────────────────────────────────────────
# COMMAND PARSERS
# ─────────────────────────────────────────────────────────────
def _parse_stoch_cmd(msg):
    parts = msg.strip().split()
    if len(parts) != 4: return None
    try:
        k, s, d = int(parts[1]), int(parts[2]), int(parts[3])
        return (k, s, d) if all(1 <= v <= 100 for v in (k, s, d)) else None
    except ValueError: return None

def _parse_set_cmd(msg):
    parts = msg.strip().split()
    if len(parts) != 4: return None
    _, tf, key, val = parts
    tf = tf.lower(); key = key.lower()
    if tf not in _TFS or key not in ('tp', 'sl'): return None
    try:
        value = int(val); return (tf, key, value) if value >= 1 else None
    except ValueError: return None

# ─────────────────────────────────────────────────────────────
# TELEGRAM UPDATE HANDLER
# ─────────────────────────────────────────────────────────────
async def process_tg_update(update: dict) -> None:
    if 'message' in update and 'text' in update['message']:
        msg = update['message']['text'].strip()
        bot_state['chat_id'] = update['message']['chat']['id']

        if msg == '/start':
            await send_tg_msg(
                '<b>Gold Scalper Bot v4.6 (MA Gap Filter Edition)</b>\n'
                'Strategies: STOCH-NEW | STOCH-OLD\n'
                'Data Engine: Deriv WebSockets (1089)\n'
                'Daily DD protection: 3%',
                get_main_keyboard()
            )

        elif msg.lower().startswith('/stoch'):
            result = _parse_stoch_cmd(msg)
            if result:
                k, s, d = result
                bot_state['stoch_k'] = k; bot_state['stoch_smooth'] = s; bot_state['stoch_d'] = d
                await send_tg_msg(f'Stoch updated: K={k} S={s} D={d}')
            else:
                await send_tg_msg('Usage: /stoch K S D  e.g. /stoch 14 3 3')

        elif msg.lower().startswith('/set'):
            result = _parse_set_cmd(msg)
            if result:
                tf, key, value = result
                bot_state['tp_pips' if key == 'tp' else 'sl_pips'][tf] = value
                tp = bot_state['tp_pips'][tf]; sl = bot_state['sl_pips'][tf]
                rr = round(tp / sl, 2) if sl else 0
                await send_tg_msg(f'Updated [{tf}]: TP={tp}p SL={sl}p RR=1:{rr}')
            else:
                await send_tg_msg('Usage: /set TF tp|sl VALUE  e.g. /set 2m sl 100')

        elif msg.startswith('/backtest'):
            try:
                start_dt = datetime.strptime(msg.split()[1], '%Y-%m-%d').replace(tzinfo=timezone.utc)
                if bot_state['is_backtesting']:
                    await send_tg_msg('A backtest is already running.', _get_cancel_kbd_for_running())
                else:
                    asyncio.create_task(run_deriv_backtest(start_dt))
            except (IndexError, ValueError):
                await send_tg_msg('Usage: /backtest YYYY-MM-DD')

        elif msg == '/cancel_bt':
            global _bt_progress
            if _bt_progress and bot_state['is_backtesting']:
                await _bt_progress.cancel(); await send_tg_msg('Cancel sent.')
            else:
                await send_tg_msg('No backtest running.')

        elif msg == '/dd_status':
            sod  = bot_state['sod_balance']
            date = bot_state['sod_date'] or '-'
            trig = 'YES - BOT PAUSED' if bot_state['dd_triggered'] else 'NO'
            if sod:
                limit_val = sod * (1 - DD_LIMIT_PCT)
                await send_tg_msg(
                    f'<b>Daily DD Protection</b>\n'
                    f'Date: {date}\n'
                    f'SOD Balance: <b>${sod:.2f}</b>\n'
                    f'Max Loss (3%): <b>${sod * DD_LIMIT_PCT:.2f}</b>\n'
                    f'Stop Equity: <b>${limit_val:.2f}</b>\n'
                    f'Triggered today: {trig}'
                )
            else:
                await send_tg_msg('DD: No SOD balance yet. Connect to server first.')

        elif msg == '/debug':
            if not bot_state['account_obj']:
                await send_tg_msg('Not connected.'); return
            try:
                raw  = await fetch_deriv_candles('5m', count=100)
                df   = calculate_indicators(pd.DataFrame(raw))
                curr = df.iloc[-2]
                now  = datetime.now(timezone.utc)
                blocked = bot_state['use_danger_filter'] and is_blocked_time(now)
                gap_val = abs(curr['ema50'] - curr['ema150'])
                await send_tg_msg(
                    f'<b>Debug [5m] - Deriv Data</b>\n'
                    f'K:{curr["K"]:.2f} D:{curr["D"]:.2f}\n'
                    f'EMA15:{curr["ema15"]:.2f} EMA50:{curr["ema50"]:.2f} EMA150:{curr["ema150"]:.2f}\n'
                    f'MA Gap: {gap_val:.2f}\n'
                    f'ATR:{curr["atr"]:.2f}\n'
                    f'Time blocked now: {"YES" if blocked else "NO"}'
                )
            except Exception as e:
                await send_tg_msg(f'Error: {e}')
        return

    if 'callback_query' not in update: return

    q       = update['callback_query']
    d       = q['data']
    chat_id = q['message']['chat']['id']
    msg_id  = q['message']['message_id']
    bot_state['chat_id'] = chat_id
    c_log(f'CB: {d}')

    try:
        await _handle_callback(d, chat_id, msg_id)
    except Exception as e:
        c_log(f'CB error [{d}]: {e}')
    finally:
        await answer_callback(q['id'])

# ─────────────────────────────────────────────────────────────
# CALLBACK HANDLER
# ─────────────────────────────────────────────────────────────
async def _handle_callback(d: str, chat_id: int, msg_id: int) -> None:
    global _bt_progress

    if d == 'noop': pass

    elif d == 'menu_main':     await edit_tg_msg(chat_id, msg_id, 'Main Menu:',     get_main_keyboard())
    elif d == 'menu_filters':  await edit_tg_msg(chat_id, msg_id, 'Filters:',       get_filters_keyboard())
    elif d == 'menu_stoch_settings': await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'menu_ma_gap':   await edit_tg_msg(chat_id, msg_id, 'MA Gap Settings:', get_ma_gap_keyboard())
    elif d == 'menu_tfs':      await edit_tg_msg(chat_id, msg_id, 'Timeframes:',    get_tf_keyboard())
    elif d == 'menu_settings': await edit_tg_msg(chat_id, msg_id, 'Risk Settings:', get_settings_keyboard())
    elif d == 'menu_backtest':
        if bot_state['is_backtesting']:
            await edit_tg_msg(chat_id, msg_id, 'Backtest running...', get_backtest_keyboard())
        else:
            await edit_tg_msg(chat_id, msg_id, f'Backtest — {_strat_label()}', get_backtest_keyboard())

    elif d == 'dd_status':
        sod = bot_state['sod_balance']; date = bot_state['sod_date'] or '-'
        trig = 'TRIGGERED' if bot_state['dd_triggered'] else 'OK'
        if sod:
            limit_val = sod * (1 - DD_LIMIT_PCT)
            text = (f'<b>Daily DD Protection (3%)</b>\n'
                    f'Date: {date}\n'
                    f'SOD: <b>${sod:.2f}</b>\n'
                    f'Max Loss: <b>${sod * DD_LIMIT_PCT:.2f}</b>\n'
                    f'Stop Equity: <b>${limit_val:.2f}</b>\n'
                    f'Status: {trig}\n\n'
                    f'Blocked hours (Damascus):\n'
                    f'  13:00-13:59 | 18:00-18:59\n'
                    f'  21:00-21:59 | 22:00-22:59')
        else:
            text = ('DD: No SOD balance yet.\n'
                    'Blocked hours (Damascus):\n'
                    '  13:00-13:59 | 18:00-18:59\n'
                    '  21:00-21:59 | 22:00-22:59')
        await edit_tg_msg(chat_id, msg_id, text, get_main_keyboard())

    elif d == 'bt_show_progress':
        await send_tg_msg(f'BT phase: {_bt_progress.phase}' if _bt_progress else 'No BT running.')

    elif d == 'cancel_bt':
        if _bt_progress and bot_state['is_backtesting']:
            await _bt_progress.cancel()
            await edit_tg_msg(chat_id, msg_id, 'Stopping backtest...', get_main_keyboard())
        else:
            await edit_tg_msg(chat_id, msg_id, 'No backtest running.', get_main_keyboard())

    elif d == 'toggle_status':
        bot_state['status'] = 'PAUSED' if bot_state['status'] == 'RUNNING' else 'RUNNING'
        if bot_state['status'] == 'RUNNING':
            if bot_state['dd_triggered']:
                bot_state['dd_triggered'] = False
                await send_tg_msg('تم استئناف البوت. تم تصفير محفز خسارة 3%.')
            
            for tf, is_active in bot_state['active_tfs'].items():
                if is_active: bot_state['cache'][tf] = pd.DataFrame()
            await send_tg_msg('🔄 <b>تم استئناف العمل</b>\nتم تفريغ الذاكرة المؤقتة. سيبدأ البوت عملية إحماء جديدة لضمان دقة البيانات 100%.')
            
        await edit_tg_msg(chat_id, msg_id, 'Main Menu:', get_main_keyboard())

    elif d == 'cycle_strategy':
        modes = ['STOCH_NEW', 'STOCH_OLD']
        bot_state['strategy_mode'] = modes[(modes.index(bot_state['strategy_mode']) + 1) % len(modes)]
        await edit_tg_msg(chat_id, msg_id, f'Strategy: {_strat_label()}', get_main_keyboard())

    elif d == 'toggle_live_conn':
        if not bot_state['live_connected']:
            await edit_tg_msg(chat_id, msg_id, 'Connecting...', get_main_keyboard())
            try:
                api = MetaApi(METAAPI_TOKEN)
                bot_state['account_obj']    = await api.metatrader_account_api.get_account(ACCOUNT_ID)
                bot_state['connection_obj'] = bot_state['account_obj'].get_rpc_connection()
                await bot_state['connection_obj'].connect()
                await bot_state['connection_obj'].wait_synchronized()
                bot_state['live_connected'] = True
                
                for tf in bot_state['timeframes']: bot_state['cache'][tf] = pd.DataFrame()
                
                await _capture_sod_balance()
                await edit_tg_msg(chat_id, msg_id, 'Connected!', get_main_keyboard())
            except Exception as e:
                await edit_tg_msg(chat_id, msg_id, f'Connection failed:\n{e}', get_main_keyboard())
        else:
            bot_state['live_connected'] = False
            bot_state['connection_obj'] = None; bot_state['account_obj'] = None
            await edit_tg_msg(chat_id, msg_id, 'Disconnected.', get_main_keyboard())

    elif d == 'set_filter_full':   bot_state['filter_mode'] = 'FULL';   await edit_tg_msg(chat_id, msg_id, 'Filters:', get_filters_keyboard())
    elif d == 'set_filter_simple': bot_state['filter_mode'] = 'SIMPLE'; await edit_tg_msg(chat_id, msg_id, 'Filters:', get_filters_keyboard())
    elif d == 'set_filter_noma':   bot_state['filter_mode'] = 'NO_MA';  await edit_tg_msg(chat_id, msg_id, 'Filters:', get_filters_keyboard())

    elif d == 'toggle_stoch_deep': bot_state['use_stoch_deep'] = not bot_state['use_stoch_deep']; await edit_tg_msg(chat_id, msg_id, 'Filters:', get_filters_keyboard())
    elif d == 'toggle_stoch_mid':  bot_state['use_stoch_mid']  = not bot_state['use_stoch_mid'];  await edit_tg_msg(chat_id, msg_id, 'Filters:', get_filters_keyboard())
    elif d == 'toggle_stoch_shal': bot_state['use_stoch_shal'] = not bot_state['use_stoch_shal']; await edit_tg_msg(chat_id, msg_id, 'Filters:', get_filters_keyboard())
    
    # --- MA Gap Callbacks ---
    elif d == 'toggle_ma_gap_fixed':
        bot_state['use_ma_gap_fixed'] = not bot_state['use_ma_gap_fixed']
        await edit_tg_msg(chat_id, msg_id, 'MA Gap Settings:', get_ma_gap_keyboard())
    elif d == 'toggle_ma_gap_atr':
        bot_state['use_ma_gap_atr'] = not bot_state['use_ma_gap_atr']
        await edit_tg_msg(chat_id, msg_id, 'MA Gap Settings:', get_ma_gap_keyboard())
    elif d == 'inc_gap_pips':
        bot_state['ma_gap_pips'] += 1
        await edit_tg_msg(chat_id, msg_id, 'MA Gap Settings:', get_ma_gap_keyboard())
    elif d == 'dec_gap_pips':
        bot_state['ma_gap_pips'] = max(1, bot_state['ma_gap_pips'] - 1)
        await edit_tg_msg(chat_id, msg_id, 'MA Gap Settings:', get_ma_gap_keyboard())
    elif d == 'inc_gap_atr':
        bot_state['ma_gap_atr_mult'] = round(bot_state['ma_gap_atr_mult'] + 0.1, 1)
        await edit_tg_msg(chat_id, msg_id, 'MA Gap Settings:', get_ma_gap_keyboard())
    elif d == 'dec_gap_atr':
        bot_state['ma_gap_atr_mult'] = max(0.1, round(bot_state['ma_gap_atr_mult'] - 0.1, 1))
        await edit_tg_msg(chat_id, msg_id, 'MA Gap Settings:', get_ma_gap_keyboard())
        
    elif d == 'toggle_danger':     bot_state['use_danger_filter'] = not bot_state['use_danger_filter']; await edit_tg_msg(chat_id, msg_id, 'Filters:', get_filters_keyboard())

    elif d == 'inc_stoch_k': bot_state['stoch_k'] = min(bot_state['stoch_k'] + 1, 100); await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'dec_stoch_k': bot_state['stoch_k'] = max(bot_state['stoch_k'] - 1, 1);   await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'inc_stoch_s': bot_state['stoch_smooth'] = min(bot_state['stoch_smooth'] + 1, 100); await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'dec_stoch_s': bot_state['stoch_smooth'] = max(bot_state['stoch_smooth'] - 1, 1);   await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'inc_stoch_d': bot_state['stoch_d'] = min(bot_state['stoch_d'] + 1, 100); await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'dec_stoch_d': bot_state['stoch_d'] = max(bot_state['stoch_d'] - 1, 1);   await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())

    elif d == 'preset_5_5_5':   bot_state['stoch_k'] = bot_state['stoch_smooth'] = bot_state['stoch_d'] = 5;    await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'preset_14_3_3':  bot_state['stoch_k'] = 14; bot_state['stoch_smooth'] = 3; bot_state['stoch_d'] = 3; await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())
    elif d == 'preset_10_3_3':  bot_state['stoch_k'] = 10; bot_state['stoch_smooth'] = 3; bot_state['stoch_d'] = 3; await edit_tg_msg(chat_id, msg_id, 'Stochastic:', get_stoch_settings_keyboard())

    elif d.startswith('toggle_tf_'):
        tf = d[len('toggle_tf_'):]
        if tf in bot_state['active_tfs']: 
            bot_state['active_tfs'][tf] = not bot_state['active_tfs'][tf]
            if bot_state['active_tfs'][tf]:
                bot_state['cache'][tf] = pd.DataFrame()
        await edit_tg_msg(chat_id, msg_id, 'Timeframes:', get_tf_keyboard())

    elif d == 'toggle_be':     bot_state['use_be']         = not bot_state['use_be'];         await edit_tg_msg(chat_id, msg_id, 'Risk:', get_settings_keyboard())
    elif d == 'toggle_atr':    bot_state['use_atr']        = not bot_state['use_atr'];        await edit_tg_msg(chat_id, msg_id, 'Risk:', get_settings_keyboard())
    elif d == 'toggle_spread': bot_state['use_max_spread'] = not bot_state['use_max_spread']; await edit_tg_msg(chat_id, msg_id, 'Risk:', get_settings_keyboard())
    elif d == 'inc_lot': bot_state['lot_size'] = round(bot_state['lot_size'] + 0.01, 2); await edit_tg_msg(chat_id, msg_id, 'Risk:', get_settings_keyboard())
    elif d == 'dec_lot': bot_state['lot_size'] = max(0.01, round(bot_state['lot_size'] - 0.01, 2)); await edit_tg_msg(chat_id, msg_id, 'Risk:', get_settings_keyboard())

    elif d == 'view_tpsl': await edit_tg_msg(chat_id, msg_id, 'TP/SL per TF:', get_tpsl_overview_keyboard())

    elif d.startswith('tpsl_edit_'):
        tf = d[len('tpsl_edit_'):]
        if tf in _TFS: await edit_tg_msg(chat_id, msg_id, f'Edit [{tf}]:', get_tpsl_edit_keyboard(tf))

    elif (d.startswith('inc_tp5_') or d.startswith('inc_tp10_') or
          d.startswith('dec_tp5_') or d.startswith('dec_tp10_')):
        if   d.startswith('inc_tp10_'): direction, step, tf = 'inc', 10, d[len('inc_tp10_'):]
        elif d.startswith('dec_tp10_'): direction, step, tf = 'dec', 10, d[len('dec_tp10_'):]
        elif d.startswith('inc_tp5_'):  direction, step, tf = 'inc',  5, d[len('inc_tp5_'):]
        else:                           direction, step, tf = 'dec',  5, d[len('dec_tp5_'):]
        if tf in _TFS:
            c = bot_state['tp_pips'][tf]
            bot_state['tp_pips'][tf] = c + step if direction == 'inc' else max(5, c - step)
            await edit_tg_msg(chat_id, msg_id, f'Edit [{tf}]:', get_tpsl_edit_keyboard(tf))

    elif (d.startswith('inc_sl5_') or d.startswith('inc_sl10_') or
          d.startswith('dec_sl5_') or d.startswith('dec_sl10_')):
        if   d.startswith('inc_sl10_'): direction, step, tf = 'inc', 10, d[len('inc_sl10_'):]
        elif d.startswith('dec_sl10_'): direction, step, tf = 'dec', 10, d[len('dec_sl10_'):]
        elif d.startswith('inc_sl5_'):  direction, step, tf = 'inc',  5, d[len('inc_sl5_'):]
        else:                           direction, step, tf = 'dec',  5, d[len('dec_sl5_'):]
        if tf in _TFS:
            c = bot_state['sl_pips'][tf]
            bot_state['sl_pips'][tf] = c + step if direction == 'inc' else max(5, c - step)
            await edit_tg_msg(chat_id, msg_id, f'Edit [{tf}]:', get_tpsl_edit_keyboard(tf))

    elif d == 'report':
        now_utc = datetime.now(timezone.utc)
        bl = blocked_time_label(now_utc) if is_blocked_time(now_utc) and bot_state['use_danger_filter'] else 'Open for trading'
        lines = [f'<b>Market Report — {_strat_label()}</b>', f'Time: {bl}']
        for tf in bot_state['timeframes']:
            if bot_state['active_tfs'][tf]: lines.append(f'[{tf}] {bot_state["market_data"][tf]}')
        await edit_tg_msg(chat_id, msg_id, '\n'.join(lines), get_main_keyboard())

    elif d == 'account':
        if not (bot_state['live_connected'] and bot_state['connection_obj']):
            await edit_tg_msg(chat_id, msg_id, 'Not connected to server.', get_main_keyboard())
        else:
            try:
                info = await bot_state['connection_obj'].get_account_information()
                pos  = await bot_state['connection_obj'].get_positions()
                sod  = bot_state['sod_balance']
                eq   = float(info.get('equity', 0))
                dd_used = f'${sod - eq:.2f} ({round((sod-eq)/sod*100,2)}%)' if sod else '-'
                text = (
                    f'<b>Account Info</b>\n'
                    f'Balance:     ${info.get("balance","N/A")}\n'
                    f'Equity:      ${eq}\n'
                    f'Free Margin: ${info.get("freeMargin","N/A")}\n'
                    f'Open trades: {len(pos)}\n'
                    f'SOD Balance: ${sod:.2f}\n'
                    f'DD Used:     {dd_used}'
                ) if sod else (
                    f'<b>Account</b>\n'
                    f'Balance:{info.get("balance","N/A")} Equity:{eq} Trades:{len(pos)}'
                )
                await edit_tg_msg(chat_id, msg_id, text, get_main_keyboard())
            except Exception as e:
                await edit_tg_msg(chat_id, msg_id, f'Error: {e}', get_main_keyboard())

    elif d.startswith('bto_adv_'):
        if bot_state['is_backtesting']:
            await edit_tg_msg(chat_id, msg_id, 'BT already running.', get_backtest_keyboard())
        else:
            days = int(d.split('_')[2])
            asyncio.create_task(run_advanced_backtest(days=days))

    elif d.startswith('bto_'):
        if bot_state['is_backtesting']:
            await edit_tg_msg(chat_id, msg_id, 'BT already running.', get_backtest_keyboard())
        else:
            days  = int(d.split('_')[1])
            start = datetime.now(timezone.utc) - timedelta(days=days)
            asyncio.create_task(run_deriv_backtest(start))

    elif d == 'close_all':
        if not (bot_state['live_connected'] and bot_state['connection_obj']):
            await edit_tg_msg(chat_id, msg_id, 'Not connected.', get_main_keyboard())
        else:
            try:
                positions = await bot_state['connection_obj'].get_positions()
                if not positions:
                    await edit_tg_msg(chat_id, msg_id, 'No open trades.', get_main_keyboard())
                else:
                    for p in positions: await bot_state['connection_obj'].close_position(p['id'])
                    await edit_tg_msg(chat_id, msg_id, f'Closed {len(positions)} trades.', get_main_keyboard())
            except Exception as e:
                await edit_tg_msg(chat_id, msg_id, f'Error: {e}', get_main_keyboard())
    else:
        c_log(f'CB unhandled: {d!r}')

# ─────────────────────────────────────────────────────────────
# TELEGRAM POLLING
# ─────────────────────────────────────────────────────────────
async def telegram_polling_loop() -> None:
    c_log('Telegram polling started.')
    url     = f'https://api.telegram.org/bot{TG_TOKEN}/getUpdates'
    poll_to = aiohttp.ClientTimeout(total=20, sock_read=15)
    backoff = 1
    while True:
        try:
            async with get_http().get(
                url, params={'offset': bot_state['last_update_id'] + 1, 'timeout': 10}, timeout=poll_to
            ) as resp:
                if resp.status == 200:
                    backoff = 1
                    for upd in (await resp.json()).get('result', []):
                        bot_state['last_update_id'] = upd['update_id']
                        asyncio.create_task(process_tg_update(upd))
                elif resp.status == 429:
                    retry_after = int(resp.headers.get('Retry-After', 5))
                    await asyncio.sleep(retry_after)
                else:
                    await asyncio.sleep(backoff); backoff = min(backoff * 2, 30)
        except asyncio.CancelledError: raise
        except Exception as e:
            c_log(f'Polling error: {e}')
            await asyncio.sleep(backoff); backoff = min(backoff * 2, 30)

# ─────────────────────────────────────────────────────────────
# TASK SUPERVISOR
# ─────────────────────────────────────────────────────────────
async def supervised(coro_fn, *args, label: str = ''):
    while True:
        try:
            await coro_fn(*args)
        except asyncio.CancelledError: raise
        except Exception as e:
            c_log(f'Task "{label or coro_fn.__name__}" crashed: {e} — restart in 5s')
            await asyncio.sleep(5)

# ─────────────────────────────────────────────────────────────
# WEB SERVER
# ─────────────────────────────────────────────────────────────
_start_time = datetime.now(timezone.utc)

async def handle_ping(request: web.Request) -> web.Response:
    uptime = str(datetime.now(timezone.utc) - _start_time).split('.')[0]
    bt     = 'RUNNING' if bot_state['is_backtesting'] else 'idle'
    live   = 'connected' if bot_state['live_connected'] else 'disconnected'
    sod    = f'${bot_state["sod_balance"]:.2f}' if bot_state['sod_balance'] else 'N/A'
    dd     = 'TRIGGERED' if bot_state['dd_triggered'] else 'OK'
    return web.Response(
        text=(
            f'Gold Scalper Bot v4.6\n'
            f'Data Engine: Deriv WebSockets\n'
            f'Uptime: {uptime}\nLive: {live}\n'
            f'Backtest: {bt}\nSOD: {sod}\nDD: {dd}'
        ), content_type='text/plain'
    )

# ─────────────────────────────────────────────────────────────
# ENTRY POINT
# ─────────────────────────────────────────────────────────────
async def main() -> None:
    get_http()
    app    = web.Application()
    app.router.add_get('/', handle_ping)
    runner = web.AppRunner(app)
    await runner.setup()
    port   = int(os.environ.get('PORT', 10000))
    await web.TCPSite(runner, '0.0.0.0', port).start()
    c_log(f'Web server on port {port}')

    tasks = [
        asyncio.create_task(supervised(timeframe_scanner, tf, label=f'scanner_{tf}'))
        for tf in bot_state['timeframes']
    ]
    tasks += [
        asyncio.create_task(supervised(telegram_polling_loop,  label='polling')),
        asyncio.create_task(supervised(position_monitor,       label='position_monitor')),
        asyncio.create_task(supervised(daily_drawdown_monitor, label='dd_monitor')),
    ]

    try:
        await asyncio.gather(*tasks)
    finally:
        if _http and not _http.closed:
            await _http.close()
        c_log('Bot shut down.')

if __name__ == '__main__':
    asyncio.run(main())
